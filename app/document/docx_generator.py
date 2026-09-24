import logging
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from app.document.base import BaseDocumentGenerator
from app.utils.helpers import ensure_directory_exists

from app.core.exceptions import DocxGenerationException

logger = logging.getLogger(__name__)


class DocxGenerator(BaseDocumentGenerator):
    def __init__(self):
        # Save documents specifically inside generated_docs/ directory
        self.output_dir = ensure_directory_exists(Path("generated_docs"))

    async def generate(self, reflection_data: Dict[str, Any]) -> Path:
        try:
            content = reflection_data.get("content", "No content provided.")
            status = reflection_data.get("status", "PASSED")
            logger.info("Generating professional Microsoft Word document using python-docx...")

            doc = Document()
            
            # Configure layout margins (standard 1-inch margins)
            for section in doc.sections:
                section.top_margin = Inches(1)
                section.bottom_margin = Inches(1)
                section.left_margin = Inches(1)
                section.right_margin = Inches(1)
                
                # Setup professional footer with generated timestamp
                footer = section.footer
                footer_p = footer.paragraphs[0]
                timestamp = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")
                footer_p.text = f"Autonomous AI Agent Report | Generated: {timestamp}"
                footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                footer_p.style.font.name = 'Calibri'
                footer_p.style.font.size = Pt(9)
                footer_p.style.font.color.rgb = RGBColor(128, 128, 128)  # Subtle Gray

            # 1. Document Title
            title_p = doc.add_paragraph()
            title_run = title_p.add_run("Autonomous AI Agent Report")
            title_run.font.name = 'Calibri'
            title_run.font.size = Pt(26)
            title_run.font.bold = True
            title_run.font.color.rgb = RGBColor(31, 78, 121)  # Deep Navy Blue
            title_p.paragraph_format.space_before = Pt(12)
            title_p.paragraph_format.space_after = Pt(24)
            title_p.alignment = WD_ALIGN_PARAGRAPH.LEFT

            # Add metadata subtitle
            sub_p = doc.add_paragraph()
            sub_p.paragraph_format.space_after = Pt(36)
            run_status = sub_p.add_run(f"Verification Status: {status}\n")
            run_status.font.bold = True
            run_status.font.size = Pt(10)
            run_status.font.color.rgb = RGBColor(0, 128, 0) if status == "PASSED" else RGBColor(180, 100, 0)
            
            run_meta = sub_p.add_run("This report was compiled and verified by the autonomous AI reasoning system.")
            run_meta.font.italic = True
            run_meta.font.size = Pt(10)
            run_meta.font.color.rgb = RGBColor(100, 100, 100)

            # 2. Parse Markdown Content and Apply Formatting
            lines = content.split("\n")
            
            for line in lines:
                stripped = line.strip()
                if not stripped:
                    continue

                # Heading 1 (e.g., "# Heading")
                if stripped.startswith("# "):
                    p = doc.add_paragraph()
                    p.paragraph_format.space_before = Pt(18)
                    p.paragraph_format.space_after = Pt(6)
                    p.paragraph_format.keep_with_next = True
                    run = p.add_run(stripped[2:])
                    run.font.name = 'Calibri'
                    run.font.size = Pt(18)
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(31, 78, 121)  # Deep Navy Blue

                # Heading 2 (e.g., "## Subheading")
                elif stripped.startswith("## "):
                    p = doc.add_paragraph()
                    p.paragraph_format.space_before = Pt(14)
                    p.paragraph_format.space_after = Pt(4)
                    p.paragraph_format.keep_with_next = True
                    run = p.add_run(stripped[3:])
                    run.font.name = 'Calibri'
                    run.font.size = Pt(14)
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(70, 130, 180)  # Steel Blue

                # Heading 3 (e.g., "### Subheading 3")
                elif stripped.startswith("### "):
                    p = doc.add_paragraph()
                    p.paragraph_format.space_before = Pt(12)
                    p.paragraph_format.space_after = Pt(4)
                    p.paragraph_format.keep_with_next = True
                    run = p.add_run(stripped[4:])
                    run.font.name = 'Calibri'
                    run.font.size = Pt(12)
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(100, 100, 100)  # Gray

                # Unordered List Items (e.g., "* item" or "- item")
                elif stripped.startswith(("* ", "- ")):
                    p = doc.add_paragraph(style='List Bullet')
                    run = p.add_run(stripped[2:])
                    run.font.name = 'Calibri'
                    run.font.size = Pt(11)
                    p.paragraph_format.space_after = Pt(3)

                # Numbered List Items (e.g., "1. item")
                elif stripped[0].isdigit() and stripped.split(".")[0].isdigit() and stripped.find(".") != -1:
                    idx = stripped.find(".")
                    num_text = stripped[idx+1:].strip()
                    p = doc.add_paragraph(style='List Number')
                    run = p.add_run(num_text)
                    run.font.name = 'Calibri'
                    run.font.size = Pt(11)
                    p.paragraph_format.space_after = Pt(3)

                # Horizontal dividers
                elif stripped == "---":
                    p = doc.add_paragraph()
                    run = p.add_run("__________________________________________________")
                    run.font.color.rgb = RGBColor(200, 200, 200)

                # Standard Paragraph
                else:
                    p = doc.add_paragraph()
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(10)  # Paragraph spacing
                    p.paragraph_format.line_spacing = 1.15   # 1.15 Line spacing
                    
                    # Check for simple markdown bold formatting (**text**)
                    parts = stripped.split("**")
                    for i, part in enumerate(parts):
                        run = p.add_run(part)
                        run.font.name = 'Calibri'
                        run.font.size = Pt(11)
                        if i % 2 == 1:
                            run.font.bold = True

            # Generate a unique document filename
            filename = f"report_{uuid.uuid4().hex[:8]}.docx"
            file_path = self.output_dir / filename
            doc.save(str(file_path))
            
            logger.info(f"Professional Microsoft Word document saved at: {file_path}")
            return file_path
        except Exception as e:
            logger.error(f"Failed to generate DOCX document: {e}", exc_info=True)
            raise DocxGenerationException(f"Error compiling Microsoft Word document: {str(e)}") from e
